"""Financial Data Export System (Stage 97).

Provides institutional CSV and XLSX export generation, spreadsheet formula-injection
protection, multi-sheet workbook assembly, secure background job lifecycle management,
and time-limited download authorization.
"""

from __future__ import annotations

import csv
import io
import os
import tempfile
from datetime import datetime, timedelta, timezone
from typing import Any, Dict, List, Optional, Tuple
from uuid import UUID, uuid4

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from atlas.screeners.enums import (
    ExportFormat,
    ExportJobStatus,
)
from atlas.screeners.errors import (
    ScreenerExportFileUnavailableError,
    ScreenerExportJobFailedError,
    ScreenerExportNotFoundError,
    ScreenerExportQuotaExceededError,
)
from atlas.screeners.schemas import (
    ScreenerExportJobDTO,
    ScreenerExportRequestDTO,
    ScreenerResultRowDTO,
    ScreenerRunResponseDTO,
)

# Prefix characters that trigger formula execution in spreadsheet applications
DANGEROUS_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def sanitize_text_cell(value: Any) -> Any:
    """
    Mitigates spreadsheet formula injection (CSV/XLSX DDE attack).

    Untrusted string values starting with '=', '+', '-', '@', '\\t', '\\r'
    are prefixed with a single quote (') so spreadsheets interpret them as literal text.

    Numeric values (including valid negative financial amounts) are NEVER modified.
    """
    if value is None:
        return ""

    # If it is already numeric (int, float) or boolean, return as-is
    if isinstance(value, (int, float, bool)):
        return value

    s = str(value)
    # Check if this string is actually a valid number (e.g. "-12.5" or "-1000")
    # Valid numbers must not be sanitized into text strings
    try:
        float(s)
        return s
    except ValueError:
        pass

    # Check for dangerous leading character
    if s and s[0] in DANGEROUS_FORMULA_PREFIXES:
        return f"'{s}"

    return s


class ScreenerExportGenerator:
    """Generates institutional CSV and XLSX export artifacts."""

    @classmethod
    def generate_csv(
        cls,
        run_response: ScreenerRunResponseDTO,
        columns: List[str],
        include_metadata: bool = True,
    ) -> bytes:
        """
        Builds a safe, Unicode-compliant UTF-8 CSV representation.
        """
        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")

        # 1. Metadata comments header
        if include_metadata:
            writer.writerow(["# ATLAS Financial Screener Export"])
            writer.writerow([f"# Query ID: {run_response.query_id}"])
            writer.writerow([f"# Universe: {run_response.universe_id}"])
            writer.writerow([f"# Evaluated As Of: {run_response.evaluated_as_of.isoformat()}"])
            writer.writerow([f"# Total Matching Rows: {run_response.total_matching_count}"])
            writer.writerow([f"# Exported Rows: {len(run_response.rows)}"])
            writer.writerow(["# Notice: Sourced data for professional research; not an investment recommendation."])
            writer.writerow([])

        # 2. Header Row
        header_labels = cls._resolve_column_headers(columns)
        writer.writerow(header_labels)

        # 3. Data Rows
        for row in run_response.rows:
            row_values = []
            for col_id in columns:
                val = cls._extract_cell_value(row, col_id)
                sanitized = sanitize_text_cell(val)
                row_values.append(sanitized)
            writer.writerow(row_values)

        return output.getvalue().encode("utf-8-sig")

    @classmethod
    def generate_xlsx(
        cls,
        run_response: ScreenerRunResponseDTO,
        columns: List[str],
        query_metadata: Optional[Dict[str, Any]] = None,
    ) -> bytes:
        """
        Builds a professional multi-sheet Excel workbook.
        Sheets:
          1. Results (Formatted financial table)
          2. Methodology (Factor definitions & ranking logic)
          3. Sources (Provider attribution & freshness status)
          4. Configuration (Universe, filters & sort spec)
        """
        wb = openpyxl.Workbook()

        # Sheet 1: Results
        ws_results = wb.active
        ws_results.title = "Results"

        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        data_font = Font(name="Segoe UI", size=10)

        # Headers
        headers = cls._resolve_column_headers(columns)
        ws_results.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws_results.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws_results.freeze_panes = "A2"

        # Data Rows
        for row in run_response.rows:
            row_data = []
            for col_id in columns:
                val = cls._extract_cell_value(row, col_id)
                row_data.append(val)

            ws_results.append(row_data)

            # Format row cells
            current_row = ws_results.max_row
            for col_idx, col_id in enumerate(columns, start=1):
                cell = ws_results.cell(row=current_row, column=col_idx)
                cell.font = data_font

                raw_val = cell.value
                # If string, sanitize against formula injection
                if isinstance(raw_val, str):
                    cell.value = sanitize_text_cell(raw_val)
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif isinstance(raw_val, (int, float)):
                    # Right-align numeric values
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    # Format standard decimals or integers
                    if isinstance(raw_val, float):
                        cell.number_format = "#,##0.00"
                    else:
                        cell.number_format = "#,##0"

        # Adjust column widths
        for col in ws_results.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_results.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Sheet 2: Methodology
        ws_method = wb.create_sheet(title="Methodology")
        ws_method.append(["ATLAS Screener Ranking & Factor Methodology"])
        ws_method.append(["Version", "1.0.0"])
        ws_method.append(["Missing Data Policy", "Complete-case exclusion; missing values are omitted, never imputed as zero."])
        ws_method.append(["Normalization", "Z-Score: sample standard deviation; Percentiles: 0.0% to 100.0% cross-sectional."])
        ws_method.append(["Disclaimer", "Quantitative factor ranks represent mathematical relative order, not investment advice."])

        # Sheet 3: Sources
        ws_sources = wb.create_sheet(title="Sources")
        ws_sources.append(["Data Sources & Freshness Breakdown"])
        ws_sources.append(["Source Coverage"])
        for k, v in run_response.source_coverage.items():
            ws_sources.append([k, v])
        ws_sources.append([])
        ws_sources.append(["Evaluated As Of", run_response.evaluated_as_of.isoformat()])

        # Sheet 4: Configuration
        ws_config = wb.create_sheet(title="Configuration")
        ws_config.append(["Screen Configuration"])
        ws_config.append(["Query ID", run_response.query_id])
        ws_config.append(["Universe ID", run_response.universe_id])
        ws_config.append(["Applied Filters Count", run_response.applied_filters_count])

        if query_metadata:
            for k, v in query_metadata.items():
                ws_config.append([str(k), str(v)])

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @classmethod
    def _resolve_column_headers(cls, column_ids: List[str]) -> List[str]:
        labels = []
        label_map = {
            "rank": "Rank",
            "composite_score": "Composite Score",
            "percentile": "Percentile (%)",
            "ticker": "Ticker",
            "canonical_security_name": "Company Name",
            "exchange_mic": "Exchange",
            "sector": "Sector",
            "industry": "Industry",
            "market_cap": "Market Cap (USD)",
            "last_price": "Last Price (USD)",
            "volume_24h": "Volume (24h)",
            "pe_ratio": "P/E Ratio",
            "revenue_growth_yoy": "Revenue Growth (YoY)",
            "roe": "ROE (%)",
            "gross_margin": "Gross Margin (%)",
            "debt_to_equity": "Debt to Equity",
            "rsi_14": "RSI (14)",
            "beta_1y": "Beta (1Y)",
        }
        for cid in column_ids:
            labels.append(label_map.get(cid, cid.replace("_", " ").title()))
        return labels

    @classmethod
    def _extract_cell_value(cls, row: ScreenerResultRowDTO, col_id: str) -> Any:
        # Check standard row attributes first
        if col_id == "rank":
            return row.rank if row.rank is not None else ""
        if col_id == "composite_score":
            return row.composite_score if row.composite_score is not None else ""
        if col_id == "percentile":
            return row.percentile if row.percentile is not None else ""
        if col_id == "ticker":
            return row.ticker
        if col_id == "canonical_security_name":
            return row.canonical_security_name
        if col_id == "exchange_mic":
            return row.exchange_mic
        if col_id == "sector":
            return row.sector or ""
        if col_id == "industry":
            return row.industry or ""

        # Check metrics dict
        if col_id in row.metrics:
            cell = row.metrics[col_id]
            if cell.status == "AVAILABLE" and cell.raw_value is not None:
                return cell.raw_value
            elif cell.status == "NOT_MEANINGFUL":
                return "N/M"
            elif cell.status == "RESTRICTED":
                return "RESTRICTED"
            return ""

        return ""


class ScreenerExportJobManager:
    """Manages asynchronous background export jobs, quotas, storage, and retention."""

    MAX_CONCURRENT_JOBS_PER_USER = 2
    MAX_TOTAL_JOBS_PER_USER = 10
    RETENTION_HOURS = 24

    def __init__(self, storage_dir: Optional[str] = None):
        self.storage_dir = storage_dir or os.path.join(tempfile.gettempdir(), "atlas_screener_exports")
        os.makedirs(self.storage_dir, exist_ok=True)
        # In-memory job registry keyed by export_job_id
        self._jobs: Dict[UUID, ScreenerExportJobDTO] = {}
        # Mapping from export_job_id to local file path
        self._file_paths: Dict[UUID, str] = {}

    def create_job(
        self,
        owner_user_id: UUID,
        request: ScreenerExportRequestDTO,
    ) -> ScreenerExportJobDTO:
        """
        Creates and enqueues a new export job subject to user quotas.
        """
        user_jobs = [j for j in self._jobs.values() if j.owner_user_id == owner_user_id]
        active_jobs = [j for j in user_jobs if j.status in (ExportJobStatus.QUEUED, ExportJobStatus.RUNNING)]

        if len(active_jobs) >= self.MAX_CONCURRENT_JOBS_PER_USER:
            raise ScreenerExportQuotaExceededError(
                f"Maximum concurrent export jobs ({self.MAX_CONCURRENT_JOBS_PER_USER}) exceeded. Please wait for current exports to complete."
            )

        if len(user_jobs) >= self.MAX_TOTAL_JOBS_PER_USER:
            # Clean up oldest completed job
            completed_jobs = sorted(
                [j for j in user_jobs if j.status in (ExportJobStatus.COMPLETED, ExportJobStatus.FAILED, ExportJobStatus.CANCELLED)],
                key=lambda x: x.created_at,
            )
            if completed_jobs:
                self.delete_job(owner_user_id, completed_jobs[0].export_job_id)

        job_id = uuid4()
        now = datetime.now(timezone.utc)
        expires_at = now + timedelta(hours=self.RETENTION_HOURS)

        job = ScreenerExportJobDTO(
            export_job_id=job_id,
            owner_user_id=owner_user_id,
            status=ExportJobStatus.QUEUED,
            format=request.format,
            scope=request.scope,
            row_count=0,
            file_size_bytes=None,
            download_url=f"/api/v1/screeners/exports/{job_id}/download",
            created_at=now,
            expires_at=expires_at,
        )

        self._jobs[job_id] = job
        return job

    def execute_job_synchronously(
        self,
        job_id: UUID,
        run_response: ScreenerRunResponseDTO,
        columns: List[str],
    ) -> ScreenerExportJobDTO:
        """
        Executes file generation, stores the artifact securely, and marks the job COMPLETED.
        """
        job = self._jobs.get(job_id)
        if not job:
            raise ScreenerExportNotFoundError(str(job_id))

        if job.status == ExportJobStatus.CANCELLED:
            return job

        job.status = ExportJobStatus.RUNNING

        try:
            if job.format == ExportFormat.XLSX:
                data = ScreenerExportGenerator.generate_xlsx(run_response, columns)
                ext = "xlsx"
            else:  # CSV default
                data = ScreenerExportGenerator.generate_csv(run_response, columns)
                ext = "csv"

            # Write file to secure storage
            filename = f"export_{job_id}.{ext}"
            filepath = os.path.join(self.storage_dir, filename)

            with open(filepath, "wb") as f:
                f.write(data)

            self._file_paths[job_id] = filepath

            job.status = ExportJobStatus.COMPLETED
            job.row_count = len(run_response.rows)
            job.file_size_bytes = len(data)
            job.completed_at = datetime.now(timezone.utc)
            return job

        except Exception as e:
            job.status = ExportJobStatus.FAILED
            job.error_message = str(e)
            raise ScreenerExportJobFailedError(str(job_id), str(e))

    def get_job(self, owner_user_id: UUID, job_id: UUID) -> ScreenerExportJobDTO:
        job = self._jobs.get(job_id)
        if not job or job.owner_user_id != owner_user_id:
            raise ScreenerExportNotFoundError(str(job_id))

        # Check expiration
        if job.expires_at and datetime.now(timezone.utc) > job.expires_at:
            job.status = ExportJobStatus.EXPIRED

        return job

    def list_jobs(self, owner_user_id: UUID) -> List[ScreenerExportJobDTO]:
        now = datetime.now(timezone.utc)
        user_jobs = []
        for j in self._jobs.values():
            if j.owner_user_id == owner_user_id:
                if j.expires_at and now > j.expires_at and j.status == ExportJobStatus.COMPLETED:
                    j.status = ExportJobStatus.EXPIRED
                user_jobs.append(j)

        return sorted(user_jobs, key=lambda x: x.created_at, reverse=True)

    def cancel_job(self, owner_user_id: UUID, job_id: UUID) -> ScreenerExportJobDTO:
        job = self.get_job(owner_user_id, job_id)
        if job.status in (ExportJobStatus.COMPLETED, ExportJobStatus.FAILED, ExportJobStatus.EXPIRED):
            return job

        job.status = ExportJobStatus.CANCELLED
        self._cleanup_file(job_id)
        return job

    def get_file_content(self, owner_user_id: UUID, job_id: UUID) -> Tuple[bytes, str, str]:
        """
        Returns (content_bytes, media_type, filename) for authorized download.
        """
        job = self.get_job(owner_user_id, job_id)
        if job.status != ExportJobStatus.COMPLETED:
            raise ScreenerExportFileUnavailableError(str(job_id))

        filepath = self._file_paths.get(job_id)
        if not filepath or not os.path.exists(filepath):
            raise ScreenerExportFileUnavailableError(str(job_id))

        with open(filepath, "rb") as f:
            content = f.read()

        if job.format == ExportFormat.XLSX:
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            filename = f"screener_export_{job_id}.xlsx"
        else:
            media_type = "text/csv; charset=utf-8"
            filename = f"screener_export_{job_id}.csv"

        return content, media_type, filename

    def delete_job(self, owner_user_id: UUID, job_id: UUID) -> None:
        self.get_job(owner_user_id, job_id)
        self._cleanup_file(job_id)
        if job_id in self._jobs:
            del self._jobs[job_id]

    def _cleanup_file(self, job_id: UUID) -> None:
        filepath = self._file_paths.pop(job_id, None)
        if filepath and os.path.exists(filepath):
            try:
                os.remove(filepath)
            except OSError:
                pass


# Singleton export manager
export_job_manager = ScreenerExportJobManager()
