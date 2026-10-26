"""Institutional Automated Tests for Stage 97:
Saved Screens, Advanced Multi-Column Sorting, Quantitative Factor Ranking, and Exports.
"""

import io
import uuid

import pytest
from atlas.auth import AuthPrincipal
from atlas.screeners.enums import (
    ExportFormat,
    ExportJobStatus,
    ExportScope,
    RankingDirection,
    RankingMethod,
    RankingMode,
    RankingTransformation,
    SortDirection,
    SortNullOrder,
)
from atlas.screeners.errors import (
    SavedScreenerConflictError,
    SavedScreenerNameInvalidError,
    SavedScreenerNotFoundError,
    SavedScreenerQuotaExceededError,
    ScreenerRankingWeightInvalidError,
    ScreenerSortLimitExceededError,
)
from atlas.screeners.export import (
    ScreenerExportGenerator,
    ScreenerExportJobManager,
    sanitize_text_cell,
)
from atlas.screeners.ranking import FactorRankingEngine
from atlas.screeners.saved import SavedScreenerRepository
from atlas.screeners.schemas import (
    FilterCriterionDTO,
    RankingFactorWeightDTO,
    SavedScreenerCreateRequestDTO,
    SavedScreenerUpdateRequestDTO,
    ScreenerExportRequestDTO,
    ScreenerRankingSpecDTO,
    ScreenerRunRequestDTO,
    SortCriterionDTO,
    UniverseSelectionDTO,
)
from atlas.screeners.service import ScreenerService
from fastapi.testclient import TestClient

from atlas_api.auth.dependencies import get_current_user
from atlas_api.main import app

# ---------------------------------------------------------------------------
# Fixtures & Test Authentication
# ---------------------------------------------------------------------------

USER_A = uuid.UUID("11111111-1111-1111-1111-111111111111")
USER_B = uuid.UUID("22222222-2222-2222-2222-222222222222")


def override_user_a():
    return AuthPrincipal(sub=str(USER_A), email="user_a@atlas.local")


def override_user_b():
    return AuthPrincipal(sub=str(USER_B), email="user_b@atlas.local")


@pytest.fixture
def client_user_a():
    app.dependency_overrides[get_current_user] = override_user_a
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


@pytest.fixture
def client_user_b():
    app.dependency_overrides[get_current_user] = override_user_b
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


# ---------------------------------------------------------------------------
# 1. Saved Screens Persistence, CRUD & Concurrency Tests
# ---------------------------------------------------------------------------


def test_saved_screener_crud_lifecycle():
    repo = SavedScreenerRepository()
    req = SavedScreenerCreateRequestDTO(
        name="High Growth Tech Screen",
        description="Screens high growth technology leaders",
        universe=UniverseSelectionDTO(universe_id="common_equity", sectors=["Information Technology"]),
        filters=[FilterCriterionDTO(filter_id="revenue_growth_yoy", value=15.0)],
        columns=["ticker", "canonical_security_name", "market_cap", "revenue_growth_yoy"],
        sort=[SortCriterionDTO(field_id="market_cap", direction=SortDirection.DESC)],
    )

    # 1. Create
    screen = repo.create(USER_A, req)
    assert screen.saved_screen_id is not None
    assert screen.owner_user_id == USER_A
    assert screen.name == "High Growth Tech Screen"
    assert screen.revision == 1
    assert screen.created_at is not None

    # 2. Read
    fetched = repo.get(USER_A, screen.saved_screen_id)
    assert fetched.saved_screen_id == screen.saved_screen_id
    assert fetched.columns == req.columns

    # 3. Update with optimistic concurrency
    update_req = SavedScreenerUpdateRequestDTO(
        name="Updated Tech Growth Strategy",
        expected_revision=1,
        columns=["ticker", "canonical_security_name", "pe_ratio"],
    )
    updated = repo.update(USER_A, screen.saved_screen_id, update_req)
    assert updated.name == "Updated Tech Growth Strategy"
    assert updated.revision == 2

    # 4. Concurrency conflict check (stale revision 1 rejected)
    with pytest.raises(SavedScreenerConflictError):
        repo.update(USER_A, screen.saved_screen_id, update_req)

    # 5. Rename
    renamed = repo.rename(USER_A, screen.saved_screen_id, "Renamed Strategy", expected_revision=2)
    assert renamed.name == "Renamed Strategy"
    assert renamed.revision == 3

    # 6. Duplicate
    cloned = repo.duplicate(USER_A, screen.saved_screen_id)
    assert cloned.saved_screen_id != screen.saved_screen_id
    assert cloned.name == "Renamed Strategy (Copy)"
    assert cloned.revision == 1

    # 7. List & Search
    list_res = repo.list(USER_A, search="Renamed")
    assert list_res.total == 2
    assert len(list_res.items) == 2

    # 8. Delete
    repo.delete(USER_A, screen.saved_screen_id)
    with pytest.raises(SavedScreenerNotFoundError):
        repo.get(USER_A, screen.saved_screen_id)


def test_saved_screener_tenant_isolation():
    repo = SavedScreenerRepository()
    req = SavedScreenerCreateRequestDTO(name="User A Private Screen")
    screen_a = repo.create(USER_A, req)

    # User B cannot access User A's private screen
    with pytest.raises(SavedScreenerNotFoundError):
        repo.get(USER_B, screen_a.saved_screen_id)

    # User B cannot update User A's private screen
    with pytest.raises(SavedScreenerNotFoundError):
        repo.update(USER_B, screen_a.saved_screen_id, SavedScreenerUpdateRequestDTO(expected_revision=1, name="Hack"))

    # User B list does not show User A's screens
    list_b = repo.list(USER_B)
    assert list_b.total == 0


def test_saved_screener_validation_and_quotas():
    repo = SavedScreenerRepository()

    # Empty name rejected
    with pytest.raises(SavedScreenerNameInvalidError):
        repo.create(USER_A, SavedScreenerCreateRequestDTO(name="   "))

    # Name exceeding 120 chars rejected
    with pytest.raises(SavedScreenerNameInvalidError):
        repo.create(USER_A, SavedScreenerCreateRequestDTO(name="A" * 121))

    # Quota limit of 100
    repo.MAX_SCREENS_PER_USER = 2
    repo.create(USER_A, SavedScreenerCreateRequestDTO(name="Screen 1"))
    repo.create(USER_A, SavedScreenerCreateRequestDTO(name="Screen 2"))

    with pytest.raises(SavedScreenerQuotaExceededError):
        repo.create(USER_A, SavedScreenerCreateRequestDTO(name="Screen 3"))


# ---------------------------------------------------------------------------
# 2. Advanced Multi-Column Sorting Tests
# ---------------------------------------------------------------------------


def test_multi_column_sorting_and_null_ordering():
    service = ScreenerService()

    # Test sorting by sector ASC, market_cap DESC, NULLS LAST
    req = ScreenerRunRequestDTO(
        sort=[
            SortCriterionDTO(field_id="sector", direction=SortDirection.ASC, null_order=SortNullOrder.NULLS_LAST),
            SortCriterionDTO(field_id="market_cap", direction=SortDirection.DESC, null_order=SortNullOrder.NULLS_LAST),
        ],
        page_size=50,
    )
    res = service.run_screen(req)
    assert len(res.rows) > 0

    # Verify sector ordering
    sectors = [r.sector or "" for r in res.rows]
    # Check that non-empty sectors are sorted alphabetically
    non_empty_sectors = [s for s in sectors if s]
    assert non_empty_sectors == sorted(non_empty_sectors)


def test_sort_limit_exceeded():
    service = ScreenerService()
    # 6 sort keys exceeds maximum 5
    sort_keys = [
        SortCriterionDTO(field_id="market_cap"),
        SortCriterionDTO(field_id="last_price"),
        SortCriterionDTO(field_id="volume_24h"),
        SortCriterionDTO(field_id="pe_ratio"),
        SortCriterionDTO(field_id="roe"),
        SortCriterionDTO(field_id="gross_margin"),
    ]
    req = ScreenerRunRequestDTO(sort=sort_keys)
    with pytest.raises(ScreenerSortLimitExceededError):
        service.run_screen(req)


# ---------------------------------------------------------------------------
# 3. Quantitative Factor Ranking Mathematics Tests
# ---------------------------------------------------------------------------


def test_single_factor_ranking_math():
    candidates = [
        {"instrument_entity_id": uuid.uuid4(), "ticker": "AAA", "canonical_security_name": "Co A", "pe": 10.0},
        {"instrument_entity_id": uuid.uuid4(), "ticker": "BBB", "canonical_security_name": "Co B", "pe": 25.0},
        {"instrument_entity_id": uuid.uuid4(), "ticker": "CCC", "canonical_security_name": "Co C", "pe": 25.0},  # Tie
        {"instrument_entity_id": uuid.uuid4(), "ticker": "DDD", "canonical_security_name": "Co D", "pe": 50.0},
    ]

    # Mode: Single factor, LOWER_VALUE_FIRST (lower PE is better), DENSE_RANK
    spec_dense = ScreenerRankingSpecDTO(
        mode=RankingMode.SINGLE_FACTOR,
        method=RankingMethod.DENSE_RANK,
        factors=[
            RankingFactorWeightDTO(
                factor_id="pe",
                direction=RankingDirection.LOWER_VALUE_FIRST,
                transformation=RankingTransformation.PERCENTILE,
            )
        ],
    )

    res = FactorRankingEngine.evaluate(candidates, spec_dense, lambda c, fid: c.get(fid))
    assert res.eligible_count == 4
    assert res.excluded_count == 0

    # Ranks under DENSE_RANK: 10.0 -> Rank 1, 25.0 -> Rank 2, 25.0 -> Rank 2, 50.0 -> Rank 3
    ranks = {r.ticker: r.rank for r in res.results}
    assert ranks["AAA"] == 1
    assert ranks["BBB"] == 2
    assert ranks["CCC"] == 2
    assert ranks["DDD"] == 3

    # Percentiles under DENSE_RANK: Rank 1 -> 100.0%, Rank 3 (out of 4 items) -> 33.33%
    pcts = {r.ticker: r.percentile for r in res.results}
    assert pcts["AAA"] == 100.0
    assert pcts["DDD"] == pytest.approx(33.33, 0.01)

    # COMPETITION_RANK test (1224)
    spec_comp = spec_dense.model_copy(update={"method": RankingMethod.COMPETITION_RANK})
    res_comp = FactorRankingEngine.evaluate(candidates, spec_comp, lambda c, fid: c.get(fid))
    ranks_comp = {r.ticker: r.rank for r in res_comp.results}
    assert ranks_comp["AAA"] == 1
    assert ranks_comp["BBB"] == 2
    assert ranks_comp["CCC"] == 2
    assert ranks_comp["DDD"] == 4  # Skips to 4 due to tie


def test_composite_factor_ranking_math_and_weights():
    candidates = [
        {"instrument_entity_id": uuid.uuid4(), "ticker": "AAA", "name": "Co A", "growth": 30.0, "pe": 15.0},
        {"instrument_entity_id": uuid.uuid4(), "ticker": "BBB", "name": "Co B", "growth": 10.0, "pe": 10.0},
        {"instrument_entity_id": uuid.uuid4(), "ticker": "CCC", "name": "Co C", "growth": 20.0, "pe": 40.0},
        {"instrument_entity_id": uuid.uuid4(), "ticker": "DDD", "name": "Co D", "growth": None, "pe": 20.0},  # Missing growth
    ]

    spec = ScreenerRankingSpecDTO(
        mode=RankingMode.MULTI_FACTOR_COMPOSITE,
        factors=[
            RankingFactorWeightDTO(
                factor_id="growth",
                direction=RankingDirection.HIGHER_VALUE_FIRST,
                transformation=RankingTransformation.Z_SCORE,
                weight=2.0,
            ),
            RankingFactorWeightDTO(
                factor_id="pe",
                direction=RankingDirection.LOWER_VALUE_FIRST,
                transformation=RankingTransformation.Z_SCORE,
                weight=1.0,
            ),
        ],
    )

    res = FactorRankingEngine.evaluate(candidates, spec, lambda c, fid: c.get(fid))
    # DDD excluded due to missing growth under complete-case policy
    assert res.eligible_count == 3
    assert res.excluded_count == 1
    assert len(res.results) == 3

    # All composite scores must be finite numbers
    for r in res.results:
        assert r.composite_score is not None
        assert "growth" in r.factor_contributions
        assert "pe" in r.factor_contributions

    # Zero or invalid weight rejection
    invalid_spec = spec.model_copy(
        update={
            "factors": [
                RankingFactorWeightDTO(factor_id="growth", weight=0.0),
            ]
        }
    )
    with pytest.raises(ScreenerRankingWeightInvalidError):
        FactorRankingEngine.evaluate(candidates, invalid_spec, lambda c, fid: c.get(fid))


# ---------------------------------------------------------------------------
# 4. Financial Data Export & Injection Protection Tests
# ---------------------------------------------------------------------------


def test_spreadsheet_formula_injection_sanitization():
    # Dangerous formula prefixes: =, +, -, @, \t, \r
    assert sanitize_text_cell("=cmd|'/C calc'!A0") == "'=cmd|'/C calc'!A0"
    assert sanitize_text_cell("+SUM(A1:A10)") == "'+SUM(A1:A10)"
    assert sanitize_text_cell("@SUM(A1:A10)") == "'@SUM(A1:A10)"
    assert sanitize_text_cell("-2+3+cmd|' /C calc'!A0") == "'-2+3+cmd|' /C calc'!A0"

    # CRITICAL: Valid numeric negative values must NOT be prefixed or altered!
    assert sanitize_text_cell(-15.5) == -15.5
    assert sanitize_text_cell("-15.5") == "-15.5"
    assert sanitize_text_cell(-1000) == -1000
    assert sanitize_text_cell("0.0") == "0.0"

    # Regular text strings without dangerous prefixes
    assert sanitize_text_cell("Apple Inc.") == "Apple Inc."
    assert sanitize_text_cell("AAPL") == "AAPL"


def test_export_generation_csv_and_xlsx():
    service = ScreenerService()
    run_res = service.run_screen(ScreenerRunRequestDTO(page_size=5))
    columns = ["ticker", "canonical_security_name", "last_price", "market_cap"]

    # 1. CSV Generation
    csv_bytes = ScreenerExportGenerator.generate_csv(run_res, columns, include_metadata=True)
    assert len(csv_bytes) > 0
    csv_text = csv_bytes.decode("utf-8-sig")
    assert "# ATLAS Financial Screener Export" in csv_text
    assert "Ticker,Company Name,Last Price (USD),Market Cap (USD)" in csv_text

    # 2. XLSX Generation
    xlsx_bytes = ScreenerExportGenerator.generate_xlsx(run_res, columns)
    assert len(xlsx_bytes) > 0
    # Must start with PK ZIP header (XLSX container)
    assert xlsx_bytes[:2] == b"PK"

    # Open with openpyxl to verify sheets
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Results" in wb.sheetnames
    assert "Methodology" in wb.sheetnames
    assert "Sources" in wb.sheetnames
    assert "Configuration" in wb.sheetnames

    # Check Results sheet
    ws = wb["Results"]
    assert ws.cell(row=1, column=1).value == "Ticker"
    assert ws.max_row == len(run_res.rows) + 1


def test_export_job_manager_lifecycle():
    manager = ScreenerExportJobManager()
    service = ScreenerService()
    run_res = service.run_screen(ScreenerRunRequestDTO(page_size=5))
    req = ScreenerExportRequestDTO(
        query=ScreenerRunRequestDTO(page_size=5),
        format=ExportFormat.CSV,
        scope=ExportScope.CURRENT_PAGE,
    )

    # 1. Create job
    job = manager.create_job(USER_A, req)
    assert job.status == ExportJobStatus.QUEUED

    # 2. Execute job
    completed = manager.execute_job_synchronously(job.export_job_id, run_res, ["ticker", "last_price"])
    assert completed.status == ExportJobStatus.COMPLETED
    assert completed.row_count == 5
    assert completed.file_size_bytes is not None

    # 3. Download authorized file
    content, media_type, filename = manager.get_file_content(USER_A, job.export_job_id)
    assert len(content) > 0
    assert "text/csv" in media_type
    assert filename.endswith(".csv")

    # 4. Clean up
    manager.delete_job(USER_A, job.export_job_id)
    assert len(manager.list_jobs(USER_A)) == 0


# ---------------------------------------------------------------------------
# 5. FastAPI Endpoints Integration Tests
# ---------------------------------------------------------------------------


def test_fastapi_saved_screener_routes():
    app.dependency_overrides[get_current_user] = override_user_a
    client = TestClient(app)
    try:
        # 1. Create saved screen
        create_payload = {
            "name": "Institutional Momentum Screen",
            "description": "Momentum strategy across common equities",
            "universe": {"universe_id": "common_equity"},
            "filters": [{"filter_id": "market_cap", "value": 5000000000.0}],
            "columns": ["ticker", "canonical_security_name", "market_cap"],
            "sort": [{"field_id": "market_cap", "direction": "desc"}],
        }
        r = client.post("/api/v1/screeners/saved", json=create_payload)
        assert r.status_code == 200
        created = r.json()
        screen_id = created["saved_screen_id"]
        assert created["name"] == "Institutional Momentum Screen"
        assert created["revision"] == 1

        # 2. List saved screens
        r = client.get("/api/v1/screeners/saved")
        assert r.status_code == 200
        list_data = r.json()
        assert list_data["total"] >= 1
        assert any(s["saved_screen_id"] == screen_id for s in list_data["items"])

        # 3. Read saved screen
        r = client.get(f"/api/v1/screeners/saved/{screen_id}")
        assert r.status_code == 200
        assert r.json()["saved_screen_id"] == screen_id

        # 4. Rerun saved screen
        r = client.post(f"/api/v1/screeners/saved/{screen_id}/run?page=1&page_size=10")
        assert r.status_code == 200
        run_data = r.json()
        assert len(run_data["rows"]) > 0

        # 5. Duplicate saved screen
        r = client.post(f"/api/v1/screeners/saved/{screen_id}/duplicate?new_name=Cloned+Strategy")
        assert r.status_code == 200
        cloned = r.json()
        assert cloned["name"] == "Cloned Strategy"
        assert cloned["saved_screen_id"] != screen_id

        # 6. Tenant isolation: User B cannot access User A's screen
        app.dependency_overrides[get_current_user] = override_user_b
        r = client.get(f"/api/v1/screeners/saved/{screen_id}")
        assert r.status_code == 404

        # 7. Delete saved screen
        app.dependency_overrides[get_current_user] = override_user_a
        r = client.delete(f"/api/v1/screeners/saved/{screen_id}")
        assert r.status_code == 200
        r = client.get(f"/api/v1/screeners/saved/{screen_id}")
        assert r.status_code == 404
    finally:
        app.dependency_overrides.clear()


def test_fastapi_export_routes(client_user_a):
    export_payload = {
        "query": {
            "universe": {"universe_id": "common_equity"},
            "filters": [],
            "page": 1,
            "page_size": 10,
        },
        "format": "csv",
        "scope": "current_page",
        "selected_columns": ["ticker", "canonical_security_name", "last_price", "market_cap"],
    }

    # 1. Create Export
    r = client_user_a.post("/api/v1/screeners/exports", json=export_payload)
    assert r.status_code == 200
    job = r.json()
    job_id = job["export_job_id"]
    assert job["status"] == "completed"

    # 2. Get status
    r = client_user_a.get(f"/api/v1/screeners/exports/{job_id}")
    assert r.status_code == 200
    assert r.json()["status"] == "completed"

    # 3. Download CSV
    r = client_user_a.get(f"/api/v1/screeners/exports/{job_id}/download")
    assert r.status_code == 200
    assert "text/csv" in r.headers["content-type"]
    assert "Ticker" in r.text

    # 4. List exports
    r = client_user_a.get("/api/v1/screeners/exports")
    assert r.status_code == 200
    assert r.json()["total"] >= 1

